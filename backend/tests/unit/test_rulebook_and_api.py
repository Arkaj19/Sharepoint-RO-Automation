from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.config import settings
from app.main import app
from app.mapping import repository
from app.rulebook import writer
from tests.conftest import put_snapshot


def test_rulebook_sheets_and_rows(env, synthetic):
    doc, ecc, s4 = synthetic
    repository.save(doc, updated_by="test", bump=False)
    out = writer.build(["TEST"])
    wb = load_workbook(out)
    assert wb.sheetnames == ["Overview", "TEST", "Plant-wise", "Lookups", "Dictionary", "Value Maps", "Filters",
                             "Ignored Columns", "Change Log"]
    ws = wb["TEST"]
    header = [c.value for c in ws[1]]
    assert header[:12] == ["Sheet Name", "Group Name", "Field Description", "Importance", "Type", "Length",
                           "Decimal", "SAP Structure", "ECC Table", "SAP Field", "Logic Type", "Logic"]
    assert ws.max_row == 1 + len(doc.fields)
    assert ws["C2"].value == "Product Number"
    assert ws["D2"].value == "mandatory for sheet"


def test_api_bootstrap_review_apply_validate(env, synthetic, monkeypatch):
    doc, ecc, s4 = synthetic
    doc.object = "MARC"                           # the validate route only accepts MARC / MBEW
    repository.save(doc, updated_by="test", bump=False)
    put_snapshot("MARC", "ECC", ecc)
    put_snapshot("MARC", "S4", s4)
    monkeypatch.setattr(settings, "LLM_PROVIDER", "none")
    c = TestClient(app)

    assert c.post("/api/agent/run", json={"object": "MARC"}).status_code == 202
    props = c.get("/api/proposals?object=MARC&status=open").json()
    pid = props[0]["proposal_id"]
    ops = c.get(f"/api/proposals/{pid}").json()["ops"]
    r = c.post(f"/api/proposals/{pid}/decisions",
               json={"decisions": [{"op_id": o["op_id"], "decision": "accepted"} for o in ops], "decided_by": "t"})
    assert r.status_code == 200
    r = c.post(f"/api/proposals/{pid}/apply", json={"applied_by": "t"})
    assert r.status_code == 200 and r.json()["applications"][0]["mapping_version"] == 2

    assert c.get("/api/mappings/MARC").json()["mapping_version"] == 2
    assert "plant_crosswalk" in c.get("/api/mappings/MARC/yaml").text
    assert "+  - id: F-MARC-plant-scope" in c.get("/api/mappings/MARC/diff?from_version=1").text
    report = c.get("/api/validate/latest?sheet=MARC").json()
    assert report["mapping_version"] == 2
    assert {r["rule_id"]: r["status"] for r in report["rules"]}["value_transformation"] == "pass"

    assert c.post("/api/rulebook", json={}).status_code == 200
    assert c.get("/api/rulebook/latest/download").status_code == 200
    snaps = c.get("/api/snapshots").json()
    assert {(d["object"], d["side"]) for d in snaps if d["versions"]} == {("MARC", "ECC"), ("MARC", "S4")}
    assert c.get("/api/preview/MARC/S4").json()["total_rows"] == len(s4)
