"""
Rule book: the Excel workbook handed to the migration team, generated from
the YAML mappings (never edited by hand - change the mapping instead).

  data/rulebooks/RuleBook_<timestamp>_MARCv7_MBEWv3.xlsx   (last 10 kept)

It uses the layout of the team's business rule workbook
(Material-Master_Source.xlsx) so it reads as a drop-in replacement:

  <OBJECT> sheets  Sheet Name | Group Name | Field Description | Importance | Type |
                   Length | Decimal | SAP Structure | ECC Table | SAP Field |
                   Logic Type | Logic            (filters in the first row's Sheet Name)
                   + governance: Status, Confidence, Agreement with S/4, Source,
                   Rule ID, Version, Approved By/At, the workbook's own logic and
                   any conflict with it, and pending proposal changes on the same row
  Overview         versions, snapshots, inputs available, latest generation result
  Plant-wise       rules that differ by plant, one column per plant block
  Lookups          supporting tables joined in, and whether they are available
  Dictionary       SAP technical field -> extract column
  Value Maps       every crosswalk entry
  Filters          every filter
  Change Log       applied operations, showing only what changed
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core import paths
from app.generation import service as generation
from app.mapping import ops as mops
from app.mapping import render, repository
from app.mapping.model import CaseStep, MappingDoc, RouteStep
from app.proposals import service as proposals
from app.store import snapshot_store as ss
from app.transforms.context import availability

KEEP_FILES = 10
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GOV_FILL = PatternFill("solid", fgColor="374151")
DRAFT_FILL = PatternFill("solid", fgColor="FFF4CC")
CONFLICT_FILL = PatternFill("solid", fgColor="FDE2E1")
WRAP = Alignment(wrap_text=True, vertical="top")

TEAM_COLUMNS = ["Sheet Name", "Group Name", "Field Description", "Importance", "Type", "Length", "Decimal",
                "SAP Structure", "ECC Table", "SAP Field", "Logic Type", "Logic"]
GOV_COLUMNS = ["Status", "Confidence", "Agreement with S/4", "Source", "Rule ID", "Version", "Approved By",
               "Approved At", "Workbook Logic Type", "Workbook Logic", "Conflict", "Proposed Logic Type",
               "Proposed Logic", "Proposal / Op", "Proposed Confidence"]
WIDTHS = {"Logic": 70, "Workbook Logic": 55, "Proposed Logic": 55, "Field Description": 34, "Sheet Name": 34,
          "Conflict": 40, "Group Name": 18, "Reason": 55, "Before": 40, "After": 40, "Expression": 60}


def _sheet(wb: Workbook, title: str, headers: list[str], gov_from: int | None = None):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill = GOV_FILL if gov_from is not None and i > gov_from else HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 15)
    ws.freeze_panes = "D2" if gov_from else "A2"
    return ws


def _finish(ws) -> None:
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def _cell(v):
    if v is None or isinstance(v, (str, int, float)):
        return v
    if isinstance(v, dict):
        return "\n".join(f"{k}: {x}" for k, x in v.items() if x not in (None, "", [], {}))
    if isinstance(v, list):
        return "\n".join(str(x) for x in v)
    return str(v)


def _change_text(before, after) -> tuple[str | None, str | None]:
    """Only the attributes that changed, e.g. 'ecc_column: - -> Material'."""
    if isinstance(before, dict) and isinstance(after, dict):
        keys = [k for k in dict.fromkeys(list(before) + list(after)) if before.get(k) != after.get(k)]
        if not keys:
            return "(no change)", "(no change)"
        return ("\n".join(f"{k}: {before.get(k) if before.get(k) not in (None, '') else '-'}" for k in keys),
                "\n".join(f"{k}: {after.get(k) if after.get(k) not in (None, '') else '-'}" for k in keys))
    return _cell(before), _cell(after)


def _latest_agreement(obj: str) -> tuple[dict[str, float], dict | None]:
    runs = generation.list_runs(obj, limit=1)
    if not runs:
        return {}, None
    rep = generation.get_run(obj, runs[0]["run_id"]) or {}
    shadow = rep.get("shadow") or {}
    return ({x["field"]: x.get("agreement") for x in shadow.get("fields", []) if x.get("agreement") is not None},
            rep)


def _drafts(doc: MappingDoc, proposal: dict | None) -> dict[str, dict]:
    """field id -> proposed logic from the open proposal's pending/accepted ops."""
    out: dict[str, dict] = {}
    if not proposal:
        return out
    for od in proposal["ops"]:
        if od["status"] not in ("pending", "accepted") or not od.get("field_id"):
            continue
        try:
            op = mops.parse_op({k: v for k, v in od.items() if k not in ("status", "decision", "preview", "result")})
            tmp = doc.model_copy(deep=True)
            mops.apply_op(tmp, op)
            f = tmp.field(od["field_id"])
            lt, text = render.logic(tmp, f) if f else ("", "")
        except (mops.OpError, ValueError, KeyError):
            lt, text = "", od.get("reason", "")
        prev = out.get(od["field_id"])
        entry = {"type": lt, "logic": text, "ref": f"{proposal['proposal_id']} / {od['op_id']}"
                 + (" (alternative)" if od.get("alternative_group") else ""), "confidence": od.get("confidence")}
        if prev:
            entry["logic"] = prev["logic"] + "\n--- or ---\n" + entry["logic"]
            entry["ref"] = prev["ref"] + "\n" + entry["ref"]
        out[od["field_id"]] = entry
    return out


def _filters_text(doc: MappingDoc) -> str:
    lines = ["Filters:"] + [f"{render.condition(f.expr)}" + (f"  [{f.source_text}]" if f.source_text else "")
                            for f in doc.filters]
    return "\n".join(lines) if doc.filters else ""


def _object_sheet(wb: Workbook, doc: MappingDoc, include_drafts: bool) -> int:
    ws = _sheet(wb, doc.object, TEAM_COLUMNS + GOV_COLUMNS, gov_from=len(TEAM_COLUMNS))
    agreement, _ = _latest_agreement(doc.object)
    open_props = proposals.list_all(doc.object, status="open", limit=1) if include_drafts else []
    drafts = _drafts(doc, proposals.get(open_props[0]["proposal_id"]) if open_props else None)
    first = True
    for f in doc.fields:
        lt, text = render.logic(doc, f)
        ref = f.reference or {}
        wbk = ref.get("workbook") or {}
        cf = ref.get("conflict")
        p = f.provenance
        table = (f.source_ref.split(".")[0] if f.source_ref else f.ecc_source.table) or ""
        sheet_cell = ""
        if first:
            sheet_cell = f"{doc.title or doc.object}\n" + _filters_text(doc)
            first = False
        d = drafts.get(f.id) or {}
        row = [sheet_cell.strip(), f.group or wbk.get("group"), f.s4_label, "mandatory for sheet" if f.mandatory else "",
               f.data_type, f.length, f.decimal, doc.s4_sheet, table, f.s4_column, lt, text,
               p.status.capitalize(), p.confidence, agreement.get(f.id), p.source, p.rule_id, p.version,
               p.approved_by, p.approved_at, wbk.get("logic_type"), wbk.get("logic"),
               cf.get("detail") if cf else None, d.get("type"), d.get("logic"), d.get("ref"), d.get("confidence")]
        ws.append(row)
        r = ws.max_row
        if d:
            for c in range(len(TEAM_COLUMNS) + len(GOV_COLUMNS) - 3, len(row) + 1):
                ws.cell(row=r, column=c).fill = DRAFT_FILL
        if cf:
            ws.cell(row=r, column=len(TEAM_COLUMNS) + 11).fill = CONFLICT_FILL
    _finish(ws)
    return len(drafts)


def _plant_wise(wb: Workbook, docs: list[MappingDoc]) -> None:
    rows, plant_cols = [], []
    for doc in docs:
        for f in doc.fields:
            t = f.transform
            if t and isinstance(t[0], RouteStep):
                per: dict[str, list[str]] = {}
                for b in t[0].branches:
                    key = (b.name or "").split(":")[0]
                    per.setdefault(key, []).extend(f"{render.condition(c.when)} → '{c.value}'" for c in b.cases)
                rows.append((doc.object, f.s4_column, f.s4_label + " (routing)", {k: "\n".join(v) for k, v in per.items()}))
            elif len(t) == 1 and isinstance(t[0], CaseStep) and all(
                    c.steps is not None and c.when.ref and c.when.ref.split(".")[-1] in ("WERKS", "BWKEY")
                    for c in t[0].cases):
                per = {}
                for c in t[0].cases:
                    label = ", ".join(c.when.values or [c.when.value])
                    tmp = f.model_copy(deep=True)
                    tmp.transform = c.steps or []
                    per[label] = render.logic(doc, tmp)[1] or "passthrough"
                rows.append((doc.object, f.s4_column, f.s4_label, per))
            else:
                continue
            for k in rows[-1][3]:
                if k not in plant_cols:
                    plant_cols.append(k)
    ws = _sheet(wb, "Plant-wise", ["Object", "SAP Field", "Field Description"] + plant_cols)
    for obj, col, label, per in rows:
        ws.append([obj, col, label] + [per.get(k, "") for k in plant_cols])
    for i in range(4, 4 + len(plant_cols)):
        ws.column_dimensions[get_column_letter(i)].width = 45
    _finish(ws)


def build(objects: list[str] | None = None, include_drafts: bool = False) -> Path:
    objects = objects or repository.list_objects()
    docs = [repository.load(o) for o in objects]
    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"
    now = datetime.now(timezone.utc).replace(microsecond=0)
    ov.append(["GyanSys Migration Tool - Rule Book"])
    ov["A1"].font = Font(bold=True, size=14)
    ov.append([f"Generated {now.isoformat()} from the YAML mappings. Do not edit - change the mapping instead."])
    ov.append([])
    head = ["Object", "Title", "Mapping version", "Updated at", "Updated by", "ECC snapshot", "S/4 snapshot",
            "Fields", "Approved", "Migrated (not reviewed)", "With a rule or source", "Lookups", "Inputs missing",
            "Filters", "Open drafts", "Last generation", "S/4 rows reproduced", "Cell agreement"]
    ov.append(head)
    for i in range(1, len(head) + 1):
        c = ov.cell(row=4, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        ov.column_dimensions[get_column_letter(i)].width = 17

    for doc in docs:
        drafts = _object_sheet(wb, doc, include_drafts)
        avail = availability(doc)
        _, rep = _latest_agreement(doc.object)
        shadow = (rep or {}).get("shadow") or {}
        ov.append([doc.object, doc.title, doc.mapping_version, doc.updated_at, doc.updated_by,
                   ss.versions(doc.object, "ECC")["current"], ss.versions(doc.object, "S4")["current"],
                   len(doc.fields), sum(1 for f in doc.fields if f.provenance.status == "approved"),
                   sum(1 for f in doc.fields if f.provenance.status == "migrated"),
                   sum(1 for f in doc.fields if doc.has_source(f) or doc.is_derived_only(f)), len(doc.lookups),
                   ", ".join(t for t, ok in avail.items() if not ok) or "-", len(doc.filters), drafts,
                   (rep or {}).get("created_at"), (shadow.get("rows") or {}).get("theirs_reproduced_share"),
                   (shadow.get("summary") or {}).get("cell_agreement")])

    _plant_wise(wb, docs)
    lk_ws = _sheet(wb, "Lookups", ["Object", "Alias", "Table", "Join", "Columns", "Available", "Description"])
    dic_ws = _sheet(wb, "Dictionary", ["Object", "Table", "SAP Field", "Extract column"])
    vm_ws = _sheet(wb, "Value Maps", ["Object", "Map ID", "Description", "On missing", "From", "To", "Rule ID",
                                      "Status", "Confidence", "Reason", "Version"])
    flt_ws = _sheet(wb, "Filters", ["Object", "Filter ID", "Side", "Expression", "Source Text", "Rule ID", "Status",
                                    "Confidence", "Reason", "Version"])
    ign_ws = _sheet(wb, "Ignored Columns", ["Object", "Side", "Column", "Reason", "Rule ID", "Status", "Confidence",
                                            "Approved By", "Version"])
    log_ws = _sheet(wb, "Change Log", ["At", "Object", "Mapping Version", "Proposal", "Op ID", "Operation", "Gate",
                                       "Target", "Before", "After", "Decided By", "Confidence", "Reason"])
    for doc in docs:
        avail = availability(doc)
        for lk in doc.lookups:
            lk_ws.append([doc.object, lk.id, lk.table, " AND ".join(f"{k.left} = {lk.table}.{k.right}" for k in lk.keys)
                          + ("".join(f" AND {w.field} {w.op} {w.value or w.values}" for w in lk.where)),
                          ", ".join(lk.columns), "yes" if avail.get(lk.table) else "NO - rules using it report "
                          "'input missing'", lk.description])
        for table, entries in sorted(doc.dictionary.items()):
            for tech, label in sorted(entries.items()):
                dic_ws.append([doc.object, table, tech, label])
        for map_id, vm in doc.value_maps.items():
            for e in vm.entries:
                p = e.provenance
                vm_ws.append([doc.object, map_id, vm.description, vm.on_missing, e.from_, e.to, p.rule_id,
                              p.status.capitalize(), p.confidence, p.reason, p.version])
        for flt in doc.filters:
            p = flt.provenance
            flt_ws.append([doc.object, flt.id, flt.side.upper(), render.condition(flt.expr), flt.source_text,
                           p.rule_id, p.status.capitalize(), p.confidence, p.reason, p.version])
        for side in ("ecc", "s4"):
            for c in getattr(doc.ignored_columns, side):
                p = c.provenance
                ign_ws.append([doc.object, side.upper(), c.column, c.reason, p.rule_id, p.status.capitalize(),
                               p.confidence, p.approved_by, p.version])
        for r in repository.read_changelog(doc.object):
            before, after = _change_text(r.get("before"), r.get("after"))
            log_ws.append([r.get("at"), r.get("object"), r.get("mapping_version"), r.get("proposal_id"),
                           r.get("op_id"), r.get("op"), r.get("gate"), _cell(r.get("target")), before, after,
                           r.get("decided_by"), r.get("confidence"), r.get("reason")])
    for ws in (lk_ws, dic_ws, vm_ws, flt_ws, log_ws):
        _finish(ws)
    ov.append([])
    ov.append(["Status legend"])
    ov.append(["Approved", "applied from an approved proposal (Agent 1 or the Databricks / workbook import)"])
    ov.append(["Migrated", "imported from the original mapping workbook; not yet reviewed in the tool"])
    ov.append(["Draft (yellow)", "pending in an open proposal - shown in the Proposed columns"])
    ov.append(["Conflict (red)", "the business workbook describes the rule differently - needs a human decision"])

    out_dir = paths.rulebooks_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    versions = "_".join(f"{d.object}v{d.mapping_version}" for d in docs)
    out = out_dir / f"RuleBook_{now.strftime('%Y%m%dT%H%M%SZ')}_{versions}{'_drafts' if include_drafts else ''}.xlsx"
    wb.save(out)
    for old in sorted(out_dir.glob("RuleBook_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)[KEEP_FILES:]:
        old.unlink(missing_ok=True)
    return out


def list_files() -> list[dict]:
    d = paths.rulebooks_dir()
    if not d.exists():
        return []
    return [{"name": p.name, "size": p.stat().st_size,
             "created_at": datetime.fromtimestamp(p.stat().st_mtime, timezone.utc).replace(microsecond=0).isoformat()}
            for p in sorted(d.glob("RuleBook_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)]
